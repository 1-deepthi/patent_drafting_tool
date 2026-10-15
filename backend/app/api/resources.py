"""Firebase-backed resource endpoints. Documents use string IDs and explicit owner IDs."""
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from datetime import datetime, timezone
from pathlib import PurePath
from app.core.security import active_user, current_user, require_roles, is_admin
from app.core.firebase import db, bucket, FirebaseNotConfigured
from app.core.config import settings
from app.services.repository import create_doc, get_doc, list_docs, patch_doc, delete_doc, now_iso
import uuid

router = APIRouter()

def _visible(item: dict, user: dict) -> bool:
    return is_admin(user) or item.get('created_by') == user['uid'] or item.get('owner_uid') == user['uid'] or user['uid'] in item.get('member_uids', []) or item.get('assigned_agent_uid') == user['uid']

def _read(collection: str, user: dict, limit: int = 100):
    return [x for x in list_docs(collection, limit=limit) if _visible(x, user)]

def _create(collection: str, body: dict, user: dict):
    body = dict(body)
    for protected in ('id', 'created_by', 'owner_uid', 'role_codes', 'account_status', 'assigned_agent_uid', 'assigned_by'):
        body.pop(protected, None)
    body['created_by'] = user['uid']
    body['owner_uid'] = user['uid']
    return create_doc(collection, body)

def _owned_get(collection: str, item_id: str, user: dict):
    item = get_doc(collection, item_id)
    if not _visible(item, user): raise HTTPException(403, 'You do not have access to this resource.')
    return item

# Teams
@router.post('/teams', status_code=201)
def create_team(body: dict, user: dict = Depends(active_user)):
    body.setdefault('member_uids', [user['uid']]); body.setdefault('team_name', body.get('name', ''))
    return _create('teams', body, user)
@router.get('/teams')
def list_teams(user: dict = Depends(active_user)):
    return _read('teams', user)
@router.get('/teams/{team_id}')
def get_team(team_id: str, user: dict = Depends(active_user)):
    return _owned_get('teams', team_id, user)
@router.patch('/teams/{team_id}')
def update_team(team_id: str, body: dict, user: dict = Depends(active_user)):
    item = _owned_get('teams', team_id, user)
    if not is_admin(user) and item.get('owner_uid') != user['uid']: raise HTTPException(403, 'Only the team owner or an administrator can update this team.')
    for protected in ('id', 'created_by', 'owner_uid', 'member_uids'):
        body.pop(protected, None)
    return patch_doc('teams', team_id, body)
@router.post('/teams/{team_id}/members', status_code=201)
def add_team_member(team_id: str, body: dict, user: dict = Depends(active_user)):
    item = _owned_get('teams', team_id, user)
    if not is_admin(user) and item.get('owner_uid') != user['uid']: raise HTTPException(403, 'Only the team owner or an administrator can manage members.')
    uid = str(body.get('uid') or body.get('user_id') or '')
    if not uid: raise HTTPException(422, 'uid is required.')
    members = list(set(item.get('member_uids', []) + [uid]))
    return patch_doc('teams', team_id, {'member_uids': members})
@router.delete('/teams/{team_id}/members/{user_id}')
def remove_team_member(team_id: str, user_id: str, user: dict = Depends(active_user)):
    item = _owned_get('teams', team_id, user)
    if not is_admin(user) and item.get('owner_uid') != user['uid']: raise HTTPException(403, 'Only the team owner or an administrator can manage members.')
    return patch_doc('teams', team_id, {'member_uids': [x for x in item.get('member_uids', []) if x != user_id]})

# Patents and assignment workflow
@router.post('/patents', status_code=201)
def create_patent(body: dict, user: dict = Depends(active_user)):
    body.setdefault('title', 'Untitled patent application'); body.setdefault('status', 'draft'); body.setdefault('workflow_status', 'draft')
    return _create('patents', body, user)
@router.get('/patents')
def list_patents(user: dict = Depends(active_user)):
    return _read('patents', user)
@router.get('/patents/available-agents')
def available_agents(_: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    return [u for u in list_docs('users') if 'PATENT_AGENT' in u.get('role_codes', []) and u.get('account_status') == 'approved']
@router.get('/patents/{patent_id}')
def get_patent(patent_id: str, user: dict = Depends(active_user)):
    return _owned_get('patents', patent_id, user)
@router.patch('/patents/{patent_id}')
def update_patent(patent_id: str, body: dict, user: dict = Depends(active_user)):
    item = _owned_get('patents', patent_id, user)
    if item.get('workflow_status') in {'submitted', 'under_review', 'approved'} and not is_admin(user) and 'PATENT_AGENT' not in user.get('role_codes', []):
        raise HTTPException(409, 'This application is no longer editable by the innovator.')
    for protected in ('id', 'created_by', 'owner_uid', 'assigned_agent_uid', 'assigned_by', 'assignment_status', 'workflow_status', 'status'):
        body.pop(protected, None)
    return patch_doc('patents', patent_id, body)
@router.post('/patents/{patent_id}/submit')
def submit_patent(patent_id: str, user: dict = Depends(active_user)):
    item = _owned_get('patents', patent_id, user)
    if not is_admin(user) and item.get('created_by') != user['uid']: raise HTTPException(403, 'Only the owner can submit this application.')
    return patch_doc('patents', patent_id, {'status': 'submitted', 'workflow_status': 'submitted', 'submitted_at': now_iso()})
@router.post('/patents/{patent_id}/assign')
def assign_agent(patent_id: str, body: dict, admin: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    agent_uid = str(body.get('agent_uid') or body.get('agent_id') or '')
    if not agent_uid: raise HTTPException(422, 'agent_uid is required.')
    agent = db().collection('users').document(agent_uid).get()
    if not agent.exists or 'PATENT_AGENT' not in (agent.to_dict() or {}).get('role_codes', []): raise HTTPException(400, 'An existing Patent Agent uid is required.')
    return patch_doc('patents', patent_id, {'assigned_agent_uid': agent_uid, 'assigned_by': admin['uid'], 'assignment_status': 'pending', 'assigned_at': now_iso()})
@router.post('/patents/{patent_id}/reassign')
def reassign_agent(patent_id: str, body: dict, admin: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    return assign_agent(patent_id, body, admin)
@router.post('/patents/{patent_id}/auto-assign')
def auto_assign_agent(patent_id: str, admin: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    agents = [u for u in list_docs('users') if 'PATENT_AGENT' in u.get('role_codes', []) and u.get('account_status') == 'approved']
    if not agents: raise HTTPException(409, 'No approved Patent Agents are available.')
    agents.sort(key=lambda u: u.get('active_assignment_count', 0))
    return patch_doc('patents', patent_id, {'assigned_agent_uid': agents[0].get('uid') or agents[0]['id'], 'assigned_by': admin['uid'], 'assignment_status': 'pending', 'assigned_at': now_iso()})
@router.post('/patents/{patent_id}/assignment/accept')
def accept_assignment(patent_id: str, user: dict = Depends(active_user)):
    item = get_doc('patents', patent_id)
    if item.get('assigned_agent_uid') != user['uid'] and not is_admin(user): raise HTTPException(403, 'Only the assigned agent can accept.')
    return patch_doc('patents', patent_id, {'assignment_status': 'accepted', 'workflow_status': 'under_review', 'accepted_at': now_iso()})
@router.post('/patents/{patent_id}/assignment/reject')
def reject_assignment(patent_id: str, body: dict | None = None, user: dict = Depends(active_user)):
    item = get_doc('patents', patent_id)
    if item.get('assigned_agent_uid') != user['uid'] and not is_admin(user): raise HTTPException(403, 'Only the assigned agent can reject.')
    return patch_doc('patents', patent_id, {'assignment_status': 'rejected', 'assignment_rejection_reason': (body or {}).get('reason', ''), 'assigned_agent_uid': None})

# Patent subresources: basic information, questionnaire, comments, calendar, drafts, documents metadata
for resource in ('applicants', 'inventors', 'comments', 'calendar-events', 'documents', 'draft-sections', 'questionnaire-versions'):
    collection = resource.replace('-', '_')
    @router.get(f'/patents/{{patent_id}}/{resource}', name=f'list_{collection}')
    def list_subresource(patent_id: str, resource=resource, collection=collection, user: dict = Depends(active_user)):
        _owned_get('patents', patent_id, user)
        items = list_docs(collection)
        return [x for x in items if x.get('patent_id') == patent_id]
    if resource in ('applicants', 'inventors', 'comments', 'calendar-events'):
        @router.post(f'/patents/{{patent_id}}/{resource}', status_code=201, name=f'create_{collection}')
        def create_subresource(patent_id: str, body: dict, resource=resource, collection=collection, user: dict = Depends(active_user)):
            _owned_get('patents', patent_id, user)
            body = dict(body); body['patent_id'] = patent_id
            if resource == 'comments': body['author_uid'] = user['uid']
            if resource == 'calendar-events': body['created_by'] = user['uid']
            return _create(collection, body, user)

@router.get('/patents/{patent_id}/questionnaire')
def get_questionnaire(patent_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    ref = db().collection('questionnaires').document(patent_id); snap = ref.get()
    if not snap.exists: return {'patent_id': patent_id, 'status': 'draft', 'answers': {}}
    out = snap.to_dict() or {}; out['id'] = snap.id; return out
@router.patch('/patents/{patent_id}/questionnaire')
def update_questionnaire(patent_id: str, body: dict, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    payload = dict(body); payload.update({'patent_id': patent_id, 'updated_by': user['uid'], 'updated_at': now_iso()})
    db().collection('questionnaires').document(patent_id).set(payload, merge=True)
    out = db().collection('questionnaires').document(patent_id).get().to_dict() or {}; out['id'] = patent_id; return out
@router.post('/patents/{patent_id}/questionnaire/submit')
def submit_questionnaire(patent_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    ref = db().collection('questionnaires').document(patent_id); ref.set({'patent_id': patent_id, 'status': 'submitted', 'submitted_by': user['uid'], 'submitted_at': now_iso()}, merge=True)
    return ref.get().to_dict() | {'id': patent_id}
@router.get('/patents/{patent_id}/questionnaire/versions')
def questionnaire_versions(patent_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user); return [x for x in list_docs('questionnaire_versions') if x.get('patent_id') == patent_id]
@router.post('/patents/{patent_id}/draft/generate')
def generate_draft(patent_id: str, body: dict | None = None, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    section = (body or {}).get('section_type', 'summary')
    content = f"Draft placeholder for {section}. Review all facts and legal wording with a qualified patent professional."
    if settings.ai_provider.lower() == 'gemini' and settings.gemini_api_key:
        try:
            import google.generativeai as genai
            genai.configure(api_key=settings.gemini_api_key)
            content = genai.GenerativeModel(settings.gemini_model).generate_content((body or {}).get('prompt', f'Draft a preliminary patent {section}.')).text
        except Exception as exc: raise HTTPException(502, f'AI provider request failed: {exc}') from exc
    doc = _create('draft_sections', {'patent_id': patent_id, 'section_type': section, 'content': content, 'generated_by': user['uid'], 'provider': settings.ai_provider}, user)
    return {'patent_id': patent_id, 'section_type': section, 'content': content, 'version': doc['id'], 'warning': 'AI output is a draft, not legal advice.'}
@router.get('/patents/{patent_id}/draft')
def list_drafts(patent_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user); return [x for x in list_docs('draft_sections') if x.get('patent_id') == patent_id]
@router.get('/patents/{patent_id}/draft/{section_type}/versions')
def draft_versions(patent_id: str, section_type: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user); return [x for x in list_docs('draft_sections') if x.get('patent_id') == patent_id and x.get('section_type') == section_type]

# Files: validation + private Firebase Storage object, metadata in Firestore
@router.get('/patents/{patent_id}/files')
def list_files(patent_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user); return [x for x in list_docs('files') if x.get('patent_id') == patent_id]
@router.post('/patents/{patent_id}/files', status_code=201)
async def upload_file(patent_id: str, file: UploadFile = File(...), user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    name = PurePath(file.filename or 'upload.bin').name; ext = PurePath(name).suffix.lower()
    content = await file.read(settings.max_upload_size_bytes + 1)
    if not content: raise HTTPException(422, 'Uploaded file is empty.')
    if len(content) > settings.max_upload_size_bytes: raise HTTPException(413, 'File exceeds 10 MB limit.')
    if ext in settings.blocked_extensions: raise HTTPException(422, f'Files with extension {ext} are not allowed.')
    key = f'patents/{patent_id}/files/{uuid.uuid4().hex}_{name}'
    try:
        blob = bucket().blob(key); blob.upload_from_string(content, content_type=file.content_type or 'application/octet-stream')
    except FirebaseNotConfigured as exc: raise HTTPException(503, str(exc)) from exc
    return _create('files', {'patent_id': patent_id, 'original_file_name': name, 'storage_key': key, 'content_type': file.content_type, 'size_bytes': len(content), 'uploaded_by': user['uid']}, user)
@router.get('/patents/{patent_id}/files/{file_id}/download')
def download_file(patent_id: str, file_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user); meta = get_doc('files', file_id)
    if meta.get('patent_id') != patent_id: raise HTTPException(404, 'File not found.')
    try:
        blob = bucket().blob(meta['storage_key']); url = blob.generate_signed_url(version='v4', expiration=300, method='GET')
    except Exception as exc: raise HTTPException(503, f'Unable to create a download link: {exc}') from exc
    return {'download_url': url, 'expires_in_seconds': 300, 'file_name': meta.get('original_file_name')}
@router.delete('/patents/{patent_id}/files/{file_id}', status_code=204)
def delete_file(patent_id: str, file_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user); meta = get_doc('files', file_id)
    if meta.get('patent_id') != patent_id: raise HTTPException(404, 'File not found.')
    if not is_admin(user) and meta.get('uploaded_by') != user['uid']: raise HTTPException(403, 'Only the uploader or an administrator can delete this file.')
    try: bucket().blob(meta['storage_key']).delete()
    except Exception: pass
    delete_doc('files', file_id)

# Notifications, calendar, templates, reports, role administration
@router.get('/notifications')
def notifications(user: dict = Depends(active_user)):
    return [x for x in list_docs('notifications') if x.get('recipient_uid') == user['uid']]
@router.get('/notifications/unread-count')
def unread_count(user: dict = Depends(active_user)):
    return {'count': sum(1 for x in list_docs('notifications') if x.get('recipient_uid') == user['uid'] and not x.get('read'))}
@router.patch('/notifications/{notification_id}/read')
def mark_notification_read(notification_id: str, user: dict = Depends(active_user)):
    item = get_doc('notifications', notification_id)
    if item.get('recipient_uid') != user['uid'] and not is_admin(user): raise HTTPException(403, 'Not your notification.')
    return patch_doc('notifications', notification_id, {'read': True, 'read_at': now_iso()})
@router.post('/notifications/read-all')
def read_all(user: dict = Depends(active_user)):
    count = 0
    for item in list_docs('notifications'):
        if item.get('recipient_uid') == user['uid'] and not item.get('read'):
            patch_doc('notifications', item['id'], {'read': True, 'read_at': now_iso()}); count += 1
    return {'updated': count}
@router.get('/calendar/events')
def my_calendar(user: dict = Depends(active_user)):
    return [x for x in list_docs('calendar_events') if x.get('created_by') == user['uid'] or x.get('owner_uid') == user['uid']]
@router.get('/templates')
def list_templates(_: dict = Depends(active_user)):
    return list_docs('templates')
@router.get('/templates/document-types')
def template_types(_: dict = Depends(active_user)):
    return sorted(set(x.get('document_type', 'unknown') for x in list_docs('templates')))
@router.post('/templates', status_code=201)
async def upload_template(file: UploadFile = File(...), document_type: str = 'PATENT_APPLICATION', user: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    content = await file.read(settings.max_upload_size_bytes + 1)
    if not content or len(content) > settings.max_upload_size_bytes: raise HTTPException(413, 'Template must be non-empty and no larger than 10 MB.')
    key = f'templates/{document_type}/{uuid.uuid4().hex}_{PurePath(file.filename or "template.docx").name}'
    blob = bucket().blob(key); blob.upload_from_string(content, content_type=file.content_type or 'application/octet-stream')
    return _create('templates', {'document_type': document_type, 'file_name': file.filename, 'storage_key': key, 'active': True, 'uploaded_by': user['uid']}, user)
@router.post('/templates/{template_id}/activate')
def activate_template(template_id: str, user: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    target = get_doc('templates', template_id)
    for t in list_docs('templates'):
        if t.get('document_type') == target.get('document_type') and t['id'] != template_id and t.get('active'):
            patch_doc('templates', t['id'], {'active': False})
    return patch_doc('templates', template_id, {'active': True, 'activated_by': user['uid']})
@router.delete('/templates/{template_id}', status_code=204)
def delete_template(template_id: str, _: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    item = get_doc('templates', template_id)
    try: bucket().blob(item['storage_key']).delete()
    except Exception: pass
    delete_doc('templates', template_id)
@router.get('/reports/dashboard')
def dashboard(user: dict = Depends(active_user)):
    patents = _read('patents', user)
    return {'total_patents': len(patents), 'draft': sum(x.get('workflow_status', x.get('status')) == 'draft' for x in patents), 'submitted': sum(x.get('workflow_status', x.get('status')) == 'submitted' for x in patents), 'under_review': sum(x.get('workflow_status') == 'under_review' for x in patents)}
@router.get('/reports')
def report_types(_: dict = Depends(active_user)):
    return [{'report_type': 'patents', 'description': 'Patent application overview'}, {'report_type': 'assignments', 'description': 'Patent agent assignment overview'}]
@router.get('/reports/{report_type}')
def report(report_type: str, user: dict = Depends(active_user)):
    if report_type == 'patents': return _read('patents', user)
    if report_type == 'assignments': return [x for x in _read('patents', user) if x.get('assigned_agent_uid')]
    raise HTTPException(404, 'Unknown report type.')
@router.get('/rbac/users')
def all_users(_: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    return list_docs('users')
@router.post('/rbac/users/{user_id}/roles')
def assign_role(user_id: str, body: dict, admin: dict = Depends(require_roles('SUPER_ADMIN'))):
    role = str(body.get('role_code', '')).upper()
    if role not in {'SUPER_ADMIN', 'ADMIN', 'PATENT_AGENT', 'INNOVATOR'}: raise HTTPException(422, 'Unknown role code.')
    ref = db().collection('users').document(user_id)
    if not ref.get().exists: raise HTTPException(404, 'User not found.')
    user = ref.get().to_dict() or {}; roles = list(set(user.get('role_codes', []) + [role]))
    ref.set({'role_codes': roles, 'role_updated_by': admin['uid'], 'updated_at': now_iso()}, merge=True)
    return {'uid': user_id, 'role_codes': roles}
@router.delete('/rbac/users/{user_id}/roles/{role_code}')
def revoke_role(user_id: str, role_code: str, admin: dict = Depends(require_roles('SUPER_ADMIN'))):
    ref = db().collection('users').document(user_id); snap = ref.get()
    if not snap.exists: raise HTTPException(404, 'User not found.')
    roles = [r for r in (snap.to_dict() or {}).get('role_codes', []) if r != role_code.upper()]
    if not roles: roles = ['INNOVATOR']
    ref.set({'role_codes': roles, 'role_updated_by': admin['uid'], 'updated_at': now_iso()}, merge=True)
    return {'uid': user_id, 'role_codes': roles}

# Editable subresources (applicants/inventors/comments are deliberately restricted).
@router.patch('/patents/{patent_id}/{resource}/{item_id}')
def patch_patent_subresource(patent_id: str, resource: str, item_id: str, body: dict, user: dict = Depends(active_user)):
    mapping = {'applicants': 'applicants', 'inventors': 'inventors', 'comments': 'comments', 'calendar-events': 'calendar_events'}
    if resource not in mapping: raise HTTPException(404, 'Unsupported patent subresource.')
    _owned_get('patents', patent_id, user)
    item = get_doc(mapping[resource], item_id)
    if item.get('patent_id') != patent_id: raise HTTPException(404, 'Subresource not found.')
    if resource == 'comments' and not is_admin(user) and item.get('created_by') != user['uid'] and item.get('author_uid') != user['uid']:
        raise HTTPException(403, 'Only the comment author or an administrator can edit this comment.')
    for protected in ('id', 'patent_id', 'created_by', 'owner_uid', 'author_uid', 'uploaded_by'):
        body.pop(protected, None)
    return patch_doc(mapping[resource], item_id, body)

@router.delete('/patents/{patent_id}/{resource}/{item_id}', status_code=204)
def delete_patent_subresource(patent_id: str, resource: str, item_id: str, user: dict = Depends(active_user)):
    mapping = {'applicants': 'applicants', 'inventors': 'inventors', 'comments': 'comments', 'calendar-events': 'calendar_events'}
    if resource not in mapping: raise HTTPException(404, 'Unsupported patent subresource.')
    _owned_get('patents', patent_id, user)
    item = get_doc(mapping[resource], item_id)
    if item.get('patent_id') != patent_id: raise HTTPException(404, 'Subresource not found.')
    if not is_admin(user) and item.get('created_by') != user['uid'] and item.get('author_uid') != user['uid']:
        raise HTTPException(403, 'Only the creator or an administrator can delete this item.')
    delete_doc(mapping[resource], item_id)

@router.post('/patents/{patent_id}/applicants/reorder')
def reorder_applicants(patent_id: str, body: dict, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    ids = body.get('ordered_ids', [])
    if not isinstance(ids, list): raise HTTPException(422, 'ordered_ids must be a list.')
    for position, item_id in enumerate(ids):
        item = get_doc('applicants', str(item_id))
        if item.get('patent_id') != patent_id: raise HTTPException(400, 'All applicant IDs must belong to this patent.')
        patch_doc('applicants', str(item_id), {'display_order': position})
    return sorted([x for x in list_docs('applicants') if x.get('patent_id') == patent_id], key=lambda x: x.get('display_order', 0))

@router.post('/patents/{patent_id}/inventors/reorder')
def reorder_inventors(patent_id: str, body: dict, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    ids = body.get('ordered_ids', [])
    if not isinstance(ids, list): raise HTTPException(422, 'ordered_ids must be a list.')
    for position, item_id in enumerate(ids):
        item = get_doc('inventors', str(item_id))
        if item.get('patent_id') != patent_id: raise HTTPException(400, 'All inventor IDs must belong to this patent.')
        patch_doc('inventors', str(item_id), {'display_order': position})
    return sorted([x for x in list_docs('inventors') if x.get('patent_id') == patent_id], key=lambda x: x.get('display_order', 0))

# Minimal DOCX generation pipeline, stored privately in Firebase Storage.
@router.get('/patents/{patent_id}/documents/jobs')
def document_jobs(patent_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    return [x for x in list_docs('documents') if x.get('patent_id') == patent_id]

@router.get('/patents/{patent_id}/documents/jobs/{job_id}')
def document_job(patent_id: str, job_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    item = get_doc('documents', job_id)
    if item.get('patent_id') != patent_id: raise HTTPException(404, 'Document job not found.')
    return item

@router.post('/patents/{patent_id}/documents/generate', status_code=201)
def generate_document(patent_id: str, body: dict | None = None, user: dict = Depends(active_user)):
    patent = _owned_get('patents', patent_id, user)
    from docx import Document
    import io
    document_type = (body or {}).get('document_type', 'PATENT_APPLICATION')
    doc = Document()
    doc.add_heading('Patent Application Draft', 0)
    doc.add_paragraph('CONFIDENTIAL — generated draft; review before use.')
    doc.add_heading('Basic Information', level=1)
    for key, value in patent.items():
        if key not in {'id', 'created_by', 'owner_uid'} and not isinstance(value, (dict, list)):
            doc.add_paragraph(f'{key.replace("_", " ").title()}: {value}')
    questionnaire = db().collection('questionnaires').document(patent_id).get()
    if questionnaire.exists:
        doc.add_heading('Questionnaire', level=1)
        for key, value in (questionnaire.to_dict() or {}).get('answers', {}).items():
            doc.add_paragraph(f'{key}: {value}')
    buf = io.BytesIO(); doc.save(buf); data = buf.getvalue()
    key = f'patents/{patent_id}/documents/{uuid.uuid4().hex}.docx'
    try:
        blob = bucket().blob(key); blob.upload_from_string(data, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    except FirebaseNotConfigured as exc: raise HTTPException(503, str(exc)) from exc
    return _create('documents', {'patent_id': patent_id, 'document_type': document_type, 'status': 'completed', 'file_name': f'{document_type.lower()}.docx', 'storage_key': key, 'content_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'size_bytes': len(data), 'generated_by': user['uid'], 'version': 1}, user)

@router.post('/patents/{patent_id}/documents/jobs/{job_id}/retry')
def retry_document_job(patent_id: str, job_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    item = get_doc('documents', job_id)
    if item.get('patent_id') != patent_id: raise HTTPException(404, 'Document job not found.')
    return generate_document(patent_id, {'document_type': item.get('document_type', 'PATENT_APPLICATION')}, user)

@router.get('/patents/{patent_id}/documents/{document_id}/download')
def download_document(patent_id: str, document_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    item = get_doc('documents', document_id)
    if item.get('patent_id') != patent_id: raise HTTPException(404, 'Document not found.')
    try:
        url = bucket().blob(item['storage_key']).generate_signed_url(version='v4', expiration=300, method='GET')
    except Exception as exc: raise HTTPException(503, f'Unable to create a download link: {exc}') from exc
    return {'download_url': url, 'expires_in_seconds': 300, 'file_name': item.get('file_name')}

@router.get('/patents/{patent_id}/documents/{document_id}/versions')
def document_versions(patent_id: str, document_id: str, user: dict = Depends(active_user)):
    _owned_get('patents', patent_id, user)
    item = get_doc('documents', document_id)
    if item.get('patent_id') != patent_id: raise HTTPException(404, 'Document not found.')
    return [{'id': item['id'], 'version': item.get('version', 1), 'created_at': item.get('created_at'), 'file_name': item.get('file_name')}]

@router.get('/rbac/matrix')
def rbac_matrix(_: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    return {'roles': ['SUPER_ADMIN', 'ADMIN', 'PATENT_AGENT', 'INNOVATOR'], 'permissions': ['users.manage', 'patents.assign', 'reports.view'], 'note': 'Role assignment is enforced by the backend; custom permission matrix editing is not implemented in this Firebase migration.'}

@router.get('/reports/permissions/matrix')
def report_permissions_matrix(_: dict = Depends(require_roles('ADMIN', 'SUPER_ADMIN'))):
    return {'roles': ['SUPER_ADMIN', 'ADMIN', 'PATENT_AGENT', 'INNOVATOR'], 'report_types': ['patents', 'assignments'], 'note': 'Static baseline; editable report-permission matrices from the original SQL backend require a dedicated migration.'}

@router.get('/reports/{report_type}/export')
def export_report(report_type: str, user: dict = Depends(active_user)):
    from fastapi.responses import JSONResponse
    data = report(report_type, user)
    return JSONResponse(content={'report_type': report_type, 'generated_at': now_iso(), 'data': data}, headers={'Content-Disposition': f'attachment; filename="{report_type}-report.json"'})
